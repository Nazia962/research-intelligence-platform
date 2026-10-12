"""Export research notes as Markdown, Word (.docx) or PDF, and build workspace briefs."""
import html
import io
import re
from datetime import date

from . import citation_service
from .paper_model import authors_text

FREE_WATERMARK = "Generated with Research Intelligence Platform (Free plan)"
DISCLAIMER = "AI-generated content must be verified against the original papers."


# ------------------------------------------------------------ markdown parse
def _is_separator(cells):
    return bool(cells) and any(cells) and all(re.fullmatch(r":?-+:?", c) for c in cells if c)


def parse_markdown(md):
    """Very small Markdown reader: headings, bullets, numbers, quotes, tables, paragraphs."""
    lines, blocks, i = str(md or "").splitlines(), [], 0
    plain_stop = re.compile(r"^(#{1,4}\s|[-*]\s|\d+[.)]\s|\||>)")
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
        elif (m := re.match(r"^(#{1,4})\s+(.*)", s)):
            blocks.append(("h", len(m.group(1)), m.group(2)))
            i += 1
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not _is_separator(cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
        elif s.startswith(">"):
            blocks.append(("quote", s.lstrip("> ").strip()))
            i += 1
        elif (m := re.match(r"^[-*]\s+(.*)", s)):
            blocks.append(("li", m.group(1)))
            i += 1
        elif (m := re.match(r"^(\d+)[.)]\s+(.*)", s)):
            blocks.append(("ol", m.group(1), m.group(2)))
            i += 1
        elif re.fullmatch(r"-{3,}|\*{3,}", s):
            blocks.append(("hr",))
            i += 1
        else:
            para = [s]
            i += 1
            while i < len(lines) and lines[i].strip() and not plain_stop.match(lines[i].strip()):
                para.append(lines[i].strip())
                i += 1
            blocks.append(("p", " ".join(para)))
    return blocks


def _inline_clean(text):
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", str(text))
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"(?<![*\w])\*(?!\*)([^*\n]+)\*(?![*\w])", r"\1", text)
    text = re.sub(r"(?<![_\w])_(?!_)([^_\n]+)_(?![_\w])", r"\1", text)
    return text


# --------------------------------------------------------------------- DOCX
def _add_runs(paragraph, text):
    for idx, seg in enumerate(_inline_clean(text).split("**")):
        if seg:
            paragraph.add_run(seg).bold = idx % 2 == 1


def markdown_to_docx(md, title, watermark=None):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(_inline_clean(title), 0)
    for b in parse_markdown(md):
        kind = b[0]
        if kind == "h":
            doc.add_heading(_inline_clean(b[2]).replace("**", ""), min(b[1], 4))
        elif kind == "li":
            _add_runs(doc.add_paragraph(style="List Bullet"), b[1])
        elif kind == "ol":
            _add_runs(doc.add_paragraph(), f"{b[1]}. {b[2]}")
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_runs(p, b[1])
            for r in p.runs:
                r.italic = True
        elif kind == "table" and b[1]:
            cols = max(len(r) for r in b[1])
            table = doc.add_table(rows=len(b[1]), cols=cols)
            table.style = "Table Grid"
            for ri, row in enumerate(b[1]):
                for ci in range(cols):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    _add_runs(cell.paragraphs[0], row[ci] if ci < len(row) else "")
                    if ri == 0:
                        for r in cell.paragraphs[0].runs:
                            r.bold = True
            doc.add_paragraph()
        elif kind == "hr":
            doc.add_paragraph("_" * 40)
        elif kind == "p":
            _add_runs(doc.add_paragraph(), b[1])
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = watermark or DISCLAIMER
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------- PDF
_PDF_REPLACE = {"→": "->", "←": "<-", "≥": ">=", "≤": "<=", "≈": "~", "×": "x",
                "✓": "v", "✔": "v", "•": "-", "−": "-", "α": "alpha", "β": "beta",
                "λ": "lambda", "μ": "u", "≠": "!=", "∞": "inf", " ": " "}


def _pdf_text(text):
    s = _inline_clean(text)
    for a, b in _PDF_REPLACE.items():
        s = s.replace(a, b)
    s = s.encode("cp1252", "ignore").decode("cp1252")
    s = html.escape(s, quote=False)
    parts = s.split("**")
    if len(parts) % 2 == 0:       # unbalanced ** markers: drop them rather than break the PDF
        return s.replace("**", "")
    return "".join(f"<b>{p}</b>" if i % 2 else p for i, p in enumerate(parts))


def markdown_to_pdf(md, title, watermark=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["BodyText"], fontSize=10, leading=14, spaceAfter=5)
    cell = ParagraphStyle("cell", parent=body, fontSize=8, leading=10, spaceAfter=0)
    heads = {1: ParagraphStyle("h1", parent=base["Heading1"], fontSize=16, spaceBefore=10, textColor="#5b3a5e"),
             2: ParagraphStyle("h2", parent=base["Heading2"], fontSize=13, spaceBefore=9, textColor="#5b3a5e"),
             3: ParagraphStyle("h3", parent=base["Heading3"], fontSize=11, spaceBefore=7),
             4: ParagraphStyle("h4", parent=base["Heading4"], fontSize=10, spaceBefore=5)}
    quote = ParagraphStyle("quote", parent=body, leftIndent=12, textColor="#555555", fontName="Helvetica-Oblique")
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=16, bulletIndent=4, spaceAfter=3)

    width = A4[0] - 40 * mm
    story = [Paragraph(_pdf_text(title), ParagraphStyle("t", parent=base["Title"], fontSize=20, textColor="#5b3a5e")),
             HRFlowable(width="100%", color=colors.HexColor("#d9cad9")), Spacer(1, 6)]
    for b in parse_markdown(md):
        kind = b[0]
        if kind == "h":
            story.append(Paragraph(_pdf_text(b[2]), heads[min(b[1], 4)]))
        elif kind == "p":
            story.append(Paragraph(_pdf_text(b[1]), body))
        elif kind == "li":
            story.append(Paragraph(_pdf_text(b[1]), bullet, bulletText="-"))
        elif kind == "ol":
            story.append(Paragraph(_pdf_text(b[2]), bullet, bulletText=f"{b[1]}."))
        elif kind == "quote":
            story.append(Paragraph(_pdf_text(b[1]), quote))
        elif kind == "hr":
            story.append(HRFlowable(width="100%", color=colors.lightgrey))
        elif kind == "table" and b[1]:
            rows, cols = b[1], max(len(r) for r in b[1])
            if cols <= 4:
                data = [[Paragraph(f"<b>{_pdf_text(c)}</b>" if ri == 0 else _pdf_text(c), cell)
                         for c in (r + [""] * (cols - len(r)))] for ri, r in enumerate(rows)]
                t = Table(data, colWidths=[width / cols] * cols, repeatRows=1)
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cdb9cd")),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3eaf3")),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                story += [t, Spacer(1, 6)]
            else:                   # wide tables: one readable block per row
                header = rows[0]
                for r in rows[1:]:
                    story.append(Paragraph(f"<b>{_pdf_text(r[0])}</b>", heads[3]))
                    for ci in range(1, min(len(r), len(header))):
                        story.append(Paragraph(f"<b>{_pdf_text(header[ci])}:</b> {_pdf_text(r[ci])}", bullet, bulletText="-"))

    note = _pdf_text(watermark or DISCLAIMER)

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(20 * mm, 10 * mm, html.unescape(re.sub(r"<[^>]+>", "", note)))
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                      bottomMargin=18 * mm, title=title).build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ------------------------------------------------------------- content builders
def _cell(text):
    return re.sub(r"\s+", " ", str(text or "")).replace("|", "/").strip()


def comparison_markdown(data):
    """Markdown for a parsed comparison dict (see ai_service.parse_comparison)."""
    head = ["Paper", "Problem", "Method", "Dataset", "Results", "Limitations"]
    lines = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    for r in data["papers"]:
        lines.append("| " + " | ".join(_cell(r.get(k, "")) for k in ("paper", "problem", "method", "dataset", "results", "limitations")) + " |")
    for key, label in (("common_themes", "Common themes"), ("key_differences", "Key differences"),
                       ("opportunity", "Opportunity for new work")):
        if data.get(key):
            lines += ["", f"## {label}", data[key]]
    return "\n".join(lines)


_ITEM_ORDER = ["plan", "comparison", "gaps", "analysis", "note"]
_ITEM_HEADINGS = {"plan": "Project plans", "comparison": "Paper comparisons", "gaps": "Research gaps",
                  "analysis": "Paper analyses", "note": "Notes"}


def _demote(text, by=3):
    return re.sub(r"^(#{1,4})\s", lambda m: "#" * min(4, len(m.group(1)) + by) + " ", text, flags=re.M)


def build_brief(ws_name, papers, items, style="APA"):
    """Markdown research brief for a whole workspace."""
    lines = [f"*Prepared {date.today():%d %B %Y} | {len(papers)} papers | {len(items)} saved items*", "",
             f"> {DISCLAIMER}", "", "## Papers"]
    for i, p in enumerate(papers, 1):
        lines.append(f"{i}. **{p['title']}** - {authors_text(p, 3)} ({p.get('year') or 'n.d.'}). {p.get('venue') or ''}".rstrip())
    for kind in _ITEM_ORDER:
        group = [it for it in items if it["kind"] == kind]
        if group:
            lines += ["", f"## {_ITEM_HEADINGS[kind]}"]
            for it in group:
                lines += ["", f"### {it['title']}", _demote(it["content"])]
    if papers:
        lines += ["", f"## References ({style})"]
        refs = citation_service.format_many(papers, style).split("\n\n" if style != "IEEE" else "\n")
        lines += [f"- {r}" if style != "IEEE" else r for r in refs]
    return "\n".join(lines)
