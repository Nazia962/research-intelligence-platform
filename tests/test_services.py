"""Offline tests: API responses are samples shaped like each service's documented output."""
import io
import json
from datetime import date, timedelta

import pytest

from services import (ai_service, citation_service, config, db, export_service, license_service,
                      paper_service, quota, ranking, sources, workspace_service)
from services.paper_model import make_key, new_paper

# ----------------------------------------------------------------- sample data
OPENALEX = {"results": [{
    "id": "https://openalex.org/W1", "doi": "https://doi.org/10.1000/XYZ", "title": "Explainable AI for Healthcare: A Survey",
    "publication_year": 2022, "cited_by_count": 321, "type": "review",
    "open_access": {"is_oa": True},
    "primary_location": {"source": {"display_name": "Journal of AI"}, "landing_page_url": "https://example.org/p1", "pdf_url": None},
    "best_oa_location": {"pdf_url": "https://example.org/p1.pdf", "landing_page_url": None},
    "authorships": [{"author": {"display_name": "Ada Lovelace"}}, {"author": {"display_name": "Alan Turing"}}],
    "abstract_inverted_index": {"We": [0], "survey": [1], "explainable": [2], "methods.": [3]}}]}

S2 = {"data": [{"paperId": "a", "title": "Explainable AI for Healthcare: a survey", "abstract": "We survey explainable methods in medicine, with a much longer abstract here that easily exceeds eighty characters.",
                "year": 2022, "venue": "J. AI", "citationCount": 400, "authors": [{"name": "Ada Lovelace"}],
                "externalIds": {"DOI": "10.1000/xyz"}, "openAccessPdf": {"url": "https://example.org/s2.pdf"},
                "url": "https://www.semanticscholar.org/paper/a", "publicationTypes": ["Review"], "isOpenAccess": True},
               {"paperId": "b", "title": "A different paper", "abstract": None, "year": 2020, "venue": "", "citationCount": 5,
                "authors": [], "externalIds": {"ArXiv": "2001.00001"}, "openAccessPdf": None, "url": "https://x.org/b",
                "publicationTypes": None, "isOpenAccess": False}]}

ARXIV = """<?xml version="1.0"?><feed xmlns="http://www.w3.org/2005/Atom" xmlns:arxiv="http://arxiv.org/schemas/atom">
<entry><id>http://arxiv.org/abs/2301.12345v1</id><published>2023-01-29T00:00:00Z</published>
<title>Deep   Learning
 for Climate</title><summary>We propose a model for climate prediction.</summary>
<author><name>Grace Hopper</name></author><link href="http://arxiv.org/abs/2301.12345v1" rel="alternate"/>
<link title="pdf" href="http://arxiv.org/pdf/2301.12345v1" rel="related" type="application/pdf"/>
<arxiv:journal_ref>Nature Climate 2023</arxiv:journal_ref></entry></feed>"""

CROSSREF = {"message": {"items": [{"DOI": "10.1000/xyz", "title": ["Explainable AI for Healthcare: A Survey"],
            "author": [{"given": "Ada", "family": "Lovelace"}], "issued": {"date-parts": [[2022, 5]]},
            "container-title": ["Journal of AI"], "is-referenced-by-count": 300,
            "abstract": "<jats:p>Short.</jats:p>", "URL": "https://doi.org/10.1000/xyz", "type": "journal-article"}]}}

EPMC = {"resultList": {"result": [{"id": "1", "pmid": "123", "doi": "10.1/abc", "title": "Cancer imaging <i>review</i>",
        "authorString": "Smith JA, Doe B.", "pubYear": "2021", "journalTitle": "Lancet", "abstractText": "<h4>Background</h4> Text.",
        "citedByCount": 10, "isOpenAccess": "Y", "pubTypeList": {"pubType": ["Review"]},
        "fullTextUrlList": {"fullTextUrl": [{"documentStyle": "html", "url": "https://h"}, {"documentStyle": "pdf", "url": "https://p.pdf"}]}}]}}

SCHOLAR = [{"title": "Scholar Paper", "link": "https://s.org/1", "snippet": "A snippet.",
            "publication_info": {"summary": "A Smith, B Jones - Journal of X, 2020 - publisher.com"},
            "resources": [{"file_format": "PDF", "link": "https://s.org/1.pdf"}],
            "inline_links": {"cited_by": {"total": 42}}}]


# ------------------------------------------------------------------- parsers
def test_parsers_return_normalised_papers():
    oa = sources.parse_openalex(OPENALEX)[0]
    assert oa["abstract"] == "We survey explainable methods." and oa["doi"] == "10.1000/xyz"
    assert oa["pdf_url"] == "https://example.org/p1.pdf" and oa["is_survey"] and oa["is_oa"]
    s2 = sources.parse_semantic_scholar(S2)
    assert s2[0]["type"] == "review" and s2[1]["abstract"] == "" and s2[1]["arxiv_id"] == "2001.00001"
    ax = sources.parse_arxiv(ARXIV)[0]
    assert ax["title"] == "Deep Learning for Climate" and ax["pdf_url"].startswith("https://arxiv.org/pdf/")
    assert ax["arxiv_id"] == "2301.12345v1" and ax["year"] == 2023 and ax["type"] == "preprint"
    cr = sources.parse_crossref(CROSSREF)[0]
    assert cr["authors"] == ["Ada Lovelace"] and cr["year"] == 2022 and cr["abstract"] == "Short."
    ep = sources.parse_europepmc(EPMC)[0]
    assert ep["title"] == "Cancer imaging review" and ep["pdf_url"] == "https://p.pdf" and ep["type"] == "review"
    assert ep["abstract"] == "Background Text." and ep["authors"] == ["Smith JA", "Doe B"]
    from services.serpapi_service import parse_scholar
    sc = parse_scholar(SCHOLAR)[0]
    assert sc["year"] == 2020 and sc["venue"] == "Journal of X" and sc["citations"] == 42 and sc["pdf_url"]


def test_merge_dedupes_and_prefers_multi_source_papers():
    results = {"openalex": sources.parse_openalex(OPENALEX), "semantic_scholar": sources.parse_semantic_scholar(S2),
               "crossref": sources.parse_crossref(CROSSREF)}
    merged = sources.merge_results(results)
    assert len(merged) == 2                                   # survey found 3x, "different paper" once
    top = merged[0]
    assert top["title"].startswith("Explainable") and set(top["sources"]) == {"openalex", "semantic_scholar", "crossref"}
    assert top["citations"] == 400 and top["key"] == make_key("10.1000/xyz", "")
    assert len({p["key"] for p in merged}) == 2


def test_search_all_survives_a_failing_source(monkeypatch):
    def boom(q, n):
        raise RuntimeError("down")
    monkeypatch.setitem(sources.SEARCHERS, "openalex", lambda q, n: sources.parse_openalex(OPENALEX))
    monkeypatch.setitem(sources.SEARCHERS, "arxiv", boom)
    out = sources.search_all("explainable ai", ["openalex", "arxiv"])
    assert len(out["papers"]) == 1 and "arxiv" in out["errors"] and out["counts"] == {"openalex": 1, "arxiv": 0}


def test_ranking_filters_and_sorts():
    papers = sources.merge_results({"semantic_scholar": sources.parse_semantic_scholar(S2), "arxiv": sources.parse_arxiv(ARXIV)})
    assert ranking.year_bounds(papers) == (2020, 2023)
    assert [p["year"] for p in ranking.sort_papers(papers, "Newest first")] == [2023, 2022, 2020]
    assert ranking.sort_papers(papers, "Most cited")[0]["citations"] == 400
    assert ranking.sort_papers(papers, "Start here (surveys first)")[0]["is_survey"]
    assert all(p["is_oa"] for p in ranking.apply_filters(papers, oa_only=True))
    assert len(ranking.apply_filters(papers, type_choice="Preprints")) == 2     # arXiv entry + S2 record with only an arXiv id
    assert len(ranking.apply_filters(papers, type_choice="Surveys / reviews")) == 1
    assert len(ranking.apply_filters(papers, min_citations=100)) == 1
    assert len(ranking.apply_filters(papers, year_range=(2021, 2024))) == 2
    assert len(ranking.apply_filters(papers, has_abstract=True)) == 1          # only the S2 survey has a long abstract


# ------------------------------------------------------------------ citations
PAPER = new_paper(title="Explainable AI for Healthcare", authors=["Ada Lovelace", "Alan M. Turing", "Smith JA"],
                  year=2022, venue="Journal of AI", doi="10.1000/xyz", url="https://example.org/p")


def test_citation_styles():
    assert citation_service.apa(PAPER) == ("Lovelace, A., Turing, A. M., & Smith, J. A. (2022). Explainable AI for Healthcare. "
                                           "Journal of AI. https://doi.org/10.1000/xyz")
    ieee = citation_service.ieee(PAPER)
    assert ieee.startswith("A. Lovelace, A. M. Turing, and J. A. Smith, “Explainable AI for Healthcare,” Journal of AI, 2022.")
    assert "doi: 10.1000/xyz" in ieee
    assert citation_service.mla(PAPER).startswith("Lovelace, A., et al.")
    bib = citation_service.bibtex(PAPER)
    assert bib.startswith("@article{Lovelace2022explainable,") and "doi = {10.1000/xyz}" in bib
    assert "TY  - JOUR" in citation_service.ris(PAPER) and "ER  - " in citation_service.ris(PAPER)
    assert citation_service.format_many([PAPER, PAPER], "IEEE").startswith("[1] ")
    pre = new_paper(title="A & B_C", authors=["Grace Hopper"], year=None, type="preprint", arxiv_id="1.2")
    assert "n.d." in citation_service.apa(pre) and "\\&" in citation_service.bibtex(pre) and "eprint" in citation_service.bibtex(pre)


# --------------------------------------------------------------- paper text
def _make_pdf(text_lines):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in text_lines:
        c.drawString(40, y, line)
        y -= 16
    c.save()
    return buf.getvalue()


LONG = ["Abstract " + "We propose a new method for detecting intrusions in networks using explainable models. " * 3,
        "Introduction", "Our dataset contains many network flows. " * 6, "However, a limitation of this work is small data. " * 3]


def test_pdf_upload_and_attach():
    data = _make_pdf(LONG)
    paper, msg = paper_service.make_uploaded_paper("my_paper.pdf", data)
    assert paper and paper["uploaded"] and "characters" in msg
    assert paper_service.extract_abstract_from_text(db.fulltext_get(paper["key"])).startswith("We propose")
    assert paper_service.get_full_text(paper)["source"] == "uploaded PDF"
    other = new_paper(title="Another paper")
    ok, _ = paper_service.attach_pdf_to_paper(other, data)
    assert ok and paper_service.get_full_text(other)["success"]
    empty, msg2 = paper_service.make_uploaded_paper("x.pdf", _make_pdf(["hi"]))
    assert empty is None and "OCR" in msg2


def test_fetch_content_html_pdf_and_failures(monkeypatch):
    html_body = ("<html><nav>menu</nav><body><p>" + "Real paper sentence about science. " * 30 + "</p></body></html>").encode()
    pdf_body = _make_pdf(LONG)
    monkeypatch.setattr(paper_service, "_is_safe_url", lambda u: True)
    monkeypatch.setattr(paper_service, "_download", lambda u: (html_body, "text/html") if u.endswith("html") else (pdf_body, "application/pdf") if u.endswith("pdf") else (b"<html>tiny</html>", "text/html"))
    assert paper_service.fetch_paper_content("https://x/a.html")["success"]
    assert "menu" not in paper_service.fetch_paper_content("https://x/a.html")["text"]
    assert paper_service.fetch_paper_content("https://x/a.pdf")["success"]
    assert not paper_service.fetch_paper_content("https://x/a.tiny")["success"]
    paper = new_paper(title="Order test", pdf_url="https://x/a.tiny", url="https://x/a.html")
    res = paper_service.get_full_text(paper)                       # PDF link fails, page works
    assert res["success"] and res["source"] == "publisher page"
    assert not paper_service.fetch_paper_content("")["success"]


def test_ssrf_guard_blocks_private_addresses():
    for url in ("http://127.0.0.1/x", "http://localhost/x", "file:///etc/passwd", "ftp://example.com/x", "http://169.254.169.254/"):
        assert not paper_service._is_safe_url(url)


# --------------------------------------------------------------------- AI
def test_ai_offline_fallback_when_no_keys():
    p = new_paper(title="T", abstract="We propose a novel approach for fraud detection using graphs. "
                  "However, a key limitation is the small dataset used for evaluation.")
    res = ai_service.analyze_paper(p, "🟢 Beginner Mode", "")
    assert res.source == "offline" and "No AI key" in res.note and "Extractive summary" in res.text
    assert "novel approach" in res.text and "small dataset" in res.text
    assert ai_service.detect_research_gaps(p, "Technical", "").source == "offline"
    assert ai_service.plan_project([p], "", {"weeks": 4, "team": 1}, "Beginner").source == "offline"
    assert ai_service.compare_papers([p, p], {}, "Beginner").source == "offline"
    assert not ai_service.analyze_paper(p, "Beginner", "", live_allowed=False).ok
    assert ai_service.analyze_paper(p, "Beginner", "", live_allowed=False).source == "limit"


def test_ai_chain_cache_and_fallback(monkeypatch):
    calls = []

    def fake_gemini(model, prompt):
        calls.append(model)
        if model == "m1":
            raise ai_service.AIError("quota", "429")
        if model == "m2":
            raise ai_service.AIError("notfound", "404")
        return "ANSWER from " + model

    monkeypatch.setattr(ai_service, "_chain", lambda: [("gemini", "m1"), ("gemini", "m2"), ("gemini", "m3")])
    monkeypatch.setattr(ai_service, "_call_gemini", fake_gemini)
    ai_service._dead_models.clear(); ai_service._cooldown.clear()
    p = new_paper(title="Chain test paper", abstract="x" * 100)
    r1 = ai_service.analyze_paper(p, "Beginner", "")
    assert r1.source == "live" and r1.model == "gemini:m3" and calls == ["m1", "m2", "m3"]
    r2 = ai_service.analyze_paper(p, "Beginner", "")                  # identical request -> cache, no API call
    assert r2.source == "cache" and r2.text == r1.text and len(calls) == 3
    r3 = ai_service.analyze_paper(p, "Technical", "")                 # different mode -> new call, skipping dead/cooling models
    assert r3.source == "live" and calls[3:] == ["m3"]


def test_ai_all_quota_gives_offline_with_message(monkeypatch):
    monkeypatch.setattr(ai_service, "_chain", lambda: [("gemini", "q1")])
    monkeypatch.setattr(ai_service, "_call_gemini", lambda m, p: (_ for _ in ()).throw(ai_service.AIError("quota")))
    ai_service._dead_models.clear(); ai_service._cooldown.clear()
    res = ai_service.detect_research_gaps(new_paper(title="Quota paper"), "Beginner", "")
    assert res.source == "offline" and "quota" in res.note.lower()


def test_mode_normalisation_and_comparison_parser():
    assert ai_service.norm_mode("🟢 Beginner Mode") == "Beginner" and ai_service.norm_mode("🔵 Technical Mode") == "Technical"
    good = '```json\n{"papers":[{"paper":"A","problem":"p","method":"m"}],"common_themes":"c"}\n```'
    parsed = ai_service.parse_comparison(good)
    assert parsed["papers"][0]["dataset"] == ai_service.NOT_AVAILABLE and parsed["common_themes"] == "c"
    assert ai_service.parse_comparison("not json") is None and ai_service.parse_comparison('{"papers": []}') is None
    off = ai_service.parse_comparison(ai_service.offline_comparison([new_paper(title="X", abstract="We propose a method. Results show improved accuracy over baselines.")]))
    assert off and len(off["papers"]) == 1


# -------------------------------------------------------- export / workspace
MD = "## Heading\nSome **bold** text with *italic* and `code`.\n\n- bullet one\n- bullet two\n1. first\n> a quote\n\n| A | B |\n|---|---|\n| 1 | two \\| x |\n---\nPlain ≥ → ✓ café"
WIDE = {"papers": [{"paper": "P1", "problem": "p", "method": "m", "dataset": "d", "results": "r", "limitations": "l"}] * 2,
        "common_themes": "themes", "key_differences": "diff", "opportunity": "opp"}


def test_export_docx_and_pdf_roundtrip():
    from docx import Document
    from pypdf import PdfReader
    docx = Document(io.BytesIO(export_service.markdown_to_docx(MD, "My Title", "wm text")))
    text = "\n".join(p.text for p in docx.paragraphs)
    assert "bold" in text and "bullet one" in text and docx.tables and docx.sections[0].footer.paragraphs[0].text == "wm text"
    pdf = PdfReader(io.BytesIO(export_service.markdown_to_pdf(MD, "My Title", "wm text")))
    page_text = pdf.pages[0].extract_text()
    assert "My Title" in page_text and "bullet one" in page_text and "wm text" in page_text
    md = export_service.comparison_markdown(WIDE)
    assert md.count("|") > 10 and "## Common themes" in md
    assert len(PdfReader(io.BytesIO(export_service.markdown_to_pdf(md, "Compare"))).pages) >= 1   # wide table path
    export_service.markdown_to_pdf("**unbalanced bold\n\n" + "word " * 5000, "Long")                 # must not raise
    assert export_service.parse_markdown("| a |\n|---|\n| b |")[0][1] == [["a"], ["b"]]


def test_workspaces_are_private_and_limits_work():
    ws, err = workspace_service.create_workspace("owner1", "Thesis")
    assert ws and not err
    assert workspace_service.create_workspace("owner1", "Thesis")[0] is None
    other, _ = workspace_service.create_workspace("owner2", "Thesis")
    papers = [new_paper(title=f"Paper {i}") for i in range(5)]
    assert workspace_service.add_papers("owner1", ws, papers, max_papers=3) == (3, 0, 2)
    assert workspace_service.add_papers("owner1", ws, papers[:1]) == (0, 1, 0)
    assert workspace_service.add_papers("owner2", ws, papers) == (0, 0, 0)          # not the owner
    assert workspace_service.list_papers("owner2", ws) == []
    item = workspace_service.save_item("owner1", ws, "analysis", "A1", "## Hello")
    assert item and workspace_service.save_item("owner2", ws, "note", "x", "y") is None
    assert workspace_service.save_item("owner1", ws, "note", "empty", "  ") is None
    workspace_service.delete_item("owner2", ws, item)
    assert len(workspace_service.list_items("owner1", ws)) == 1
    backup = workspace_service.export_json("owner1", ws)
    assert workspace_service.export_json("owner2", ws) == ""
    new_id, msg = workspace_service.import_json("owner1", backup)
    assert new_id and "3 papers" in msg and len(workspace_service.list_papers("owner1", new_id)) == 3
    assert workspace_service.import_json("owner1", "{}")[0] is None
    assert workspace_service.import_json("owner1", "garbage")[0] is None
    workspace_service.delete_workspace("owner2", ws)                                   # wrong owner: no-op
    assert len(workspace_service.list_workspaces("owner1")) == 2
    workspace_service.remove_paper("owner1", ws, papers[0]["key"])
    assert len(workspace_service.list_papers("owner1", ws)) == 2
    workspace_service.delete_workspace("owner1", ws)
    assert [w["name"] for w in workspace_service.list_workspaces("owner1")] == ["Thesis (imported 1)"]
    brief = export_service.build_brief("Thesis", papers[:2], [{"kind": "analysis", "title": "A1", "content": "# Big heading\ntext"}])
    assert "### A1" in brief and "#### Big heading" in brief and "## References (APA)" in brief


# ------------------------------------------------------------ licences / quota
def test_license_roundtrip_and_plans(monkeypatch):
    monkeypatch.setattr(config, "LICENSE_SECRET", "s3cret")
    monkeypatch.setattr(config, "MONETIZATION_ENABLED", True)
    key = license_service.make_key("Nazia K!", 30)
    assert key.startswith("RIP-PRO-") and "NAZIAK" in key
    assert license_service.verify_key(key)[0] and license_service.verify_key(key.lower())[0]
    assert not license_service.verify_key(key[:-1] + ("A" if key[-1] != "A" else "B"))[0]
    assert not license_service.verify_key("RIP-PRO-20990101-NAZIAK-AAAAAAAAAAAA")[0]
    old = license_service.make_key("old", 1, today=date.today() - timedelta(days=5))
    ok, msg, _ = license_service.verify_key(old)
    assert not ok and "expired" in msg
    monkeypatch.setattr(config, "LICENSE_REVOKED", {key})
    assert "revoked" in license_service.verify_key(key)[1]
    monkeypatch.setattr(config, "LICENSE_REVOKED", set())
    assert license_service.current_plan("u1") == "free"
    assert license_service.activate("u1", key)[0] and license_service.current_plan("u1") == "pro"
    assert not license_service.activate("u2", "nonsense")[0] and license_service.current_plan("u2") == "free"
    monkeypatch.setattr(config, "MONETIZATION_ENABLED", False)
    assert license_service.current_plan("u2") == "open"


def test_quota_limits(monkeypatch):
    assert quota.limits_for("free").compare_max == 3 and quota.limits_for("pro").compare_max == 5
    assert quota.limits_for("open").watermark is False and quota.limits_for("free").watermark is True
    owner = "quota-owner"
    for _ in range(config.FREE_AI_PER_DAY):
        assert quota.can_use_ai(owner, "free")[0]
        quota.record_ai_use(owner)
    allowed, why = quota.can_use_ai(owner, "free")
    assert not allowed and "free AI" in why
    assert quota.can_use_ai(owner, "pro")[0] and quota.can_use_ai(owner, "open")[0]
    monkeypatch.setattr(config, "GLOBAL_AI_PER_DAY", 1)
    assert not quota.can_use_ai("anyone", "pro")[0]


def test_research_radar_isolates_source_errors(monkeypatch):
    from services import discovery_service
    monkeypatch.setattr(discovery_service, "search_scholar", lambda q, limit=8: [{"title": "Paper"}])
    monkeypatch.setattr(discovery_service, "search_web", lambda q: "web evidence")
    def fail_news(q):
        raise RuntimeError("news quota")
    monkeypatch.setattr(discovery_service, "search_news", fail_news)
    result = discovery_service.research_sweep("research topic")
    assert result["scholar"][0]["title"] == "Paper"
    assert result["web"] == "web evidence"
    assert "news" in result["errors"]
